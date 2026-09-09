from io import BytesIO
from html import escape
import base64
from sqlalchemy import select
from app.core.models import *
from app.core.db import SessionLocal
from app.cases.service import records, snapshot, row
from app.audit.service import audit
from app.evidence.storage import storage
from app.verification.service import report_guard

def queue_report(db,who,case):
    db.execute(select(Case).where(Case.id==case.id).with_for_update())
    report_guard(db,case)
    previous=db.scalar(select(Report).where(Report.case_id==case.id,Report.status.in_([ReportStatus.QUEUED,ReportStatus.GENERATING,ReportStatus.READY])))
    if previous: return previous,False
    report=Report(case_id=case.id); db.add(report); db.flush()
    audit(db,who,case.id,'REPORT_REQUESTED','report',report.id,{'officer':who.name,'completion_attested':True})
    db.commit()
    return report,True

def report_content(db,case):
    product=snapshot(db,case.id)
    declarations=[]; findings=[]
    for d in records(db,DeclarationResult,case_id=case.id):
        review=db.scalar(select(AuditEvent).where(AuditEvent.entity_id==d.id,AuditEvent.event_type=='DECLARATION_REVIEWED').order_by(AuditEvent.created_at.desc()))
        declarations.append({'title':GROUPS[d.declaration_group],'value':review.payload.get('selected_value') if review else d.fused_value,
            'evidence_state':str(d.evidence_state),'review':row(review)})
    for f in records(db,Finding,case_id=case.id):
        if f.status not in (FindingStatus.ACCEPTED,FindingStatus.CORRECTED): continue
        event=db.scalar(select(VerificationEvent).where(VerificationEvent.finding_id==f.id))
        if not event: continue
        rr=db.get(RuleResult,f.rule_result_id); rule=db.get(Rule,rr.rule_id); officer=db.get(User,event.officer_id)
        findings.append({'id':f.id,'status':str(f.status),'rule_code':rule.rule_code,'rule_version':rule.version,'legal_reference':rule.legal_reference,
            'explanation':rr.explanation,'value':f.verified_value,'officer':officer.name,'reason':event.reason,'time':event.created_at.isoformat(),'sources':rr.inputs_json})
    return {'case':row(case),'product':row(product),'declarations':declarations,'findings':findings,
        'evidence':[row(e) for e in records(db,EvidenceItem,case_id=case.id) if e.source_type==EvidenceSource.FIELD_CAPTURE and e.storage_key],
        'history':[row(a) for a in records(db,AuditEvent,case_id=case.id) if a.event_type in ('DECLARATION_REVIEWED','FINDING_VERIFIED','REPORT_REQUESTED')]}

def display(value):
    if value is None: return 'No value established'
    if isinstance(value,dict): return '; '.join(f'{k.replace("_"," ")}: {v}' for k,v in value.items()) or 'No declaration detected in reviewed search area'
    return str(value)

def generate_report(report_id):
    with SessionLocal() as db:
        report=db.get(Report,report_id)
        if not report or report.status!=ReportStatus.QUEUED: return
        report.status=ReportStatus.GENERATING; db.commit()
        try:
            case=db.get(Case,report.case_id); report_guard(db,case); content=report_content(db,case)
            from docx import Document
            from docx.shared import Inches, Pt, RGBColor
            from weasyprint import HTML, default_url_fetcher
            doc=Document(); section=doc.sections[0]
            section.top_margin=section.bottom_margin=Inches(0.8)
            section.left_margin=section.right_margin=Inches(0.8)
            doc.styles['Normal'].font.name='Arial'; doc.styles['Normal'].font.size=Pt(10)
            doc.add_heading('e-Maap · Inspection record',0)
            doc.add_paragraph('Measure Compliance. Protect Consumers.')
            doc.add_paragraph('SIH prototype • Source-supported evidence is distinct from officer-verified findings.')
            generated=now(); reference='EM-RPT-'+report.id[:8].upper()
            metadata=[('Report',reference),('Generated',generated.isoformat()),('Case',case.id),('Product',content['product']['product_name']),('Manufacturer',content['product']['manufacturer']),('Batch / lot',content['product']['batch_lot'] or 'Not recorded'),('Retailer',case.shop_name)]
            for key,value in metadata: doc.add_paragraph(f'{key}: {value}')
            html=['<html><head><meta charset="utf-8"><style>@page{size:A4;margin:20mm;@bottom-right{content:counter(page)}} body{font:10pt sans-serif;color:#17212b} h1,h2{color:#173b5e} h1{font-size:22pt} h2{font-size:14pt;margin-top:20pt} table{border-collapse:collapse;width:100%}td,th{border:1px solid #d8e0e7;padding:7pt;text-align:left;vertical-align:top} th{background:#f6f8fa} img{max-width:100%;max-height:80mm} figure{break-inside:avoid} small{overflow-wrap:anywhere} .note{color:#52606d}</style></head><body><h1>e-Maap · Inspection record</h1><p>Measure Compliance. Protect Consumers.</p><p class="note">SIH prototype. Source-supported evidence is distinct from officer-verified findings.</p>']
            html.extend(f'<p><b>{escape(k)}:</b> {escape(str(v))}</p>' for k,v in metadata)
            doc.add_heading('Reviewed declarations',1); html.append('<h2>Reviewed declarations</h2><table><tr><th>Declaration</th><th>Reviewed value</th><th>Evidence state</th></tr>')
            table=doc.add_table(rows=1,cols=3); table.style='Table Grid'
            for c,title in zip(table.rows[0].cells,['Declaration','Reviewed value','Evidence state']): c.text=title
            for d in content['declarations']:
                cells=table.add_row().cells
                for cell,text in zip(cells,[d['title'],display(d['value']),d['evidence_state']]): cell.text=text
                html.append(f'<tr><td>{escape(d["title"])}</td><td>{escape(display(d["value"]))}</td><td>{escape(d["evidence_state"])}</td></tr>')
            html.append('</table><h2>Officer-verified findings</h2>'); doc.add_heading('Officer-verified findings',1)
            if not content['findings']:
                notice='No officer-verified adverse finding is recorded for the supported checks. This is not certification against every Legal Metrology requirement.'
                html.append('<p>'+escape(notice)+'</p>'); doc.add_paragraph(notice)
            for f in content['findings']:
                title=f'{f["rule_code"]} · {f["status"]}'
                details=f'{f["explanation"]}\nOfficer: {f["officer"]}\nReason: {f["reason"]}\nDecision time: {f["time"]}\nRule version: {f["rule_version"]}\nReference: {f["legal_reference"]}\nFinding ID: {f["id"]}'
                doc.add_heading(title,2); doc.add_paragraph(details)
                html.append(f'<h3>{escape(title)}</h3><p>{escape(details).replace(chr(10),"<br>")}</p>')
            doc.add_heading('Evidence and provenance',1); html.append('<h2>Evidence and provenance</h2>')
            for e in content['evidence']:
                image=storage().get(e['storage_key'])
                caption=f'#{e["sequence_no"]} · {e["id"]}\nCapture: {e["captured_at"]} · Sync: {e["synced_at"]}\nGPS: {e["latitude"]}, {e["longitude"]} ± {e["accuracy_m"]} m\nSHA-256: {e["sha256"]}'
                doc.add_picture(BytesIO(image),width=Inches(5.7)); doc.add_paragraph(caption)
                html.append(f'<figure><img src="data:{e["mime_type"]};base64,{base64.b64encode(image).decode()}"/><figcaption><small>{escape(caption).replace(chr(10),"<br>")}</small></figcaption></figure>')
            doc.add_heading('Officer review history',1); html.append('<h2>Officer review history</h2>')
            for entry in content['history']:
                text=f'{entry["created_at"]} · {entry["event_type"]} · {entry["actor_id"]}\n{entry["payload"].get("reason", "Completion requested by officer")} '
                doc.add_paragraph(text); html.append('<p>'+escape(text)+'</p>')
            footer='Rule references use the frozen SRS prototype fixture version. The report is a prototype inspection record; no automatic legal decision is issued.'
            doc.add_paragraph(footer); html.append('<p class="note">'+escape(footer)+'</p></body></html>')
            def fetch(url,*args,**kwargs):
                if not url.startswith('data:'): raise ValueError('Report rendering may load embedded evidence only')
                return default_url_fetcher(url,*args,**kwargs)
            pdf=HTML(string=''.join(html),url_fetcher=fetch).write_pdf()
            editable=BytesIO(); doc.save(editable)
            report.storage_key=f'reports/{case.id}/{report.id}.pdf'; report.editable_storage_key=f'reports/{case.id}/{report.id}.docx'
            storage().put(report.storage_key,pdf,'application/pdf'); storage().put(report.editable_storage_key,editable.getvalue(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document')
            report.status=ReportStatus.READY; report.generated_at=generated; case.status=CaseStatus.COMPLETED
            assignment=db.scalar(select(Assignment).where(Assignment.case_id==case.id))
            if assignment: assignment.status=AssignmentStatus.COMPLETED; assignment.completed_at=generated
            audit(db,None,case.id,'REPORT_READY','report',report.id,{'reference':reference}); db.commit()
        except Exception as exc:
            db.rollback(); report=db.get(Report,report_id); report.status=ReportStatus.FAILED; report.failure_code='REPORT_GENERATION_FAILED'
            audit(db,None,report.case_id,'REPORT_FAILED','report',report.id,{'detail':str(exc)[:500]}); db.commit()
